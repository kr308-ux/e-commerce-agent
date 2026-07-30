"""Safe, read-only spreadsheet parsing for creator imports."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree

import openpyxl
import xlrd

from .import_types import SheetData


ALLOWED_EXTENSIONS = {".xlsx", ".xls", ".csv"}
XLSX_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/zip",
    "application/octet-stream",
}
XLS_MIME_TYPES = {
    "application/vnd.ms-excel",
    "application/octet-stream",
}
CSV_MIME_TYPES = {
    "text/csv",
    "text/plain",
    "application/csv",
    "application/vnd.ms-excel",
    "application/octet-stream",
}
ZERO_FORMAT_PATTERN = re.compile(r"^0+$")


class SpreadsheetValidationError(ValueError):
    """Raised when an uploaded spreadsheet is unsafe or unsupported."""


def _repair_invalid_xlsx_styles(file_bytes: bytes) -> bytes | None:
    """Repair empty fill nodes emitted by some third-party XLSX exporters."""

    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as source:
            styles = source.read("xl/styles.xml")
            root = ElementTree.fromstring(styles)
            repaired_fill_count = 0
            for fill in root.iter():
                if fill.tag.rsplit("}", 1)[-1] != "fill" or len(fill):
                    continue
                namespace = (
                    fill.tag.split("}", 1)[0][1:]
                    if fill.tag.startswith("{")
                    else ""
                )
                pattern_fill_tag = (
                    f"{{{namespace}}}patternFill" if namespace else "patternFill"
                )
                ElementTree.SubElement(
                    fill,
                    pattern_fill_tag,
                    {"patternType": "none"},
                )
                repaired_fill_count += 1

            if not repaired_fill_count:
                return None

            if root.tag.startswith("{"):
                ElementTree.register_namespace("", root.tag.split("}", 1)[0][1:])
            repaired_styles = ElementTree.tostring(
                root,
                encoding="utf-8",
                xml_declaration=True,
            )
            output = io.BytesIO()
            with zipfile.ZipFile(output, "w") as destination:
                for item in source.infolist():
                    content = (
                        repaired_styles
                        if item.filename == "xl/styles.xml"
                        else source.read(item.filename)
                    )
                    destination.writestr(item, content)
            return output.getvalue()
    except (ElementTree.ParseError, KeyError, OSError, zipfile.BadZipFile):
        return None


def _load_xlsx(file_bytes: bytes):
    options = {
        "read_only": True,
        "data_only": True,
        "keep_links": False,
    }
    try:
        return openpyxl.load_workbook(io.BytesIO(file_bytes), **options)
    except (TypeError, ValueError) as original_error:
        repaired = _repair_invalid_xlsx_styles(file_bytes)
        if repaired is not None:
            try:
                return openpyxl.load_workbook(io.BytesIO(repaired), **options)
            except (TypeError, ValueError, OSError, zipfile.BadZipFile):
                pass
        raise SpreadsheetValidationError(
            "XLSX 文件结构无法解析，请用 Excel 或 WPS 另存为标准 XLSX 后重试。"
        ) from original_error


def safe_file_name(file_name: str) -> str:
    name = Path(str(file_name or "")).name
    cleaned = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff._ -]+", "_", name).strip()
    return cleaned[:255] or "creator-import"


def validate_upload(
    *,
    file_name: str,
    content_type: str,
    file_bytes: bytes,
    max_size: int,
) -> str:
    if not file_bytes:
        raise SpreadsheetValidationError("上传文件为空。")
    if len(file_bytes) > max_size:
        raise SpreadsheetValidationError(
            f"文件大小不能超过 {max_size // (1024 * 1024)} MB。"
        )

    extension = Path(file_name).suffix.casefold()
    if extension not in ALLOWED_EXTENSIONS:
        raise SpreadsheetValidationError("仅支持 .xlsx、.xls 和 .csv 文件。")

    normalized_mime = str(content_type or "").split(";", 1)[0].strip().casefold()
    allowed_mimes = {
        ".xlsx": XLSX_MIME_TYPES,
        ".xls": XLS_MIME_TYPES,
        ".csv": CSV_MIME_TYPES,
    }[extension]
    if normalized_mime and normalized_mime not in allowed_mimes:
        raise SpreadsheetValidationError("文件类型与扩展名不匹配。")

    if extension == ".xlsx":
        if not file_bytes.startswith(b"PK"):
            raise SpreadsheetValidationError("文件内容不是有效的 XLSX。")
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
                names = set(archive.namelist())
                if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                    raise SpreadsheetValidationError("XLSX 文件结构不完整。")
                if any(name.casefold().endswith("vbaproject.bin") for name in names):
                    raise SpreadsheetValidationError("不允许上传包含宏的 Excel 文件。")
                total_uncompressed = sum(item.file_size for item in archive.infolist())
                if total_uncompressed > max_size * 20:
                    raise SpreadsheetValidationError("XLSX 解压后体积异常。")
        except zipfile.BadZipFile as error:
            raise SpreadsheetValidationError("XLSX 文件已损坏。") from error
    elif extension == ".xls":
        if not file_bytes.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
            raise SpreadsheetValidationError("文件内容不是有效的 XLS。")
    else:
        if b"\x00" in file_bytes[:4096]:
            raise SpreadsheetValidationError("CSV 文件包含不支持的二进制内容。")
    return extension


def _decode_csv(file_bytes: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise SpreadsheetValidationError("CSV 编码无法识别，请使用 UTF-8 或 GB18030。")


def _string_value(value: object, *, number_format: str = "") -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, int):
        normalized_format = str(number_format or "").split(";", 1)[0]
        if ZERO_FORMAT_PATTERN.fullmatch(normalized_format):
            return f"{value:0{len(normalized_format)}d}"
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            integer = int(value)
            normalized_format = str(number_format or "").split(";", 1)[0]
            if ZERO_FORMAT_PATTERN.fullmatch(normalized_format):
                return f"{integer:0{len(normalized_format)}d}"
            return str(integer)
        return format(Decimal(str(value)), "f")
    return str(value)


def list_sheet_names(file_name: str, file_bytes: bytes) -> tuple[str, ...]:
    extension = Path(file_name).suffix.casefold()
    if extension == ".xlsx":
        workbook = _load_xlsx(file_bytes)
        try:
            return tuple(workbook.sheetnames)
        finally:
            workbook.close()
    if extension == ".xls":
        workbook = xlrd.open_workbook(file_contents=file_bytes, on_demand=True)
        try:
            return tuple(workbook.sheet_names())
        finally:
            workbook.release_resources()
    return ("CSV",)


def read_sheet(
    *,
    file_name: str,
    file_bytes: bytes,
    sheet_name: str | None = None,
    max_rows: int | None = None,
) -> SheetData:
    extension = Path(file_name).suffix.casefold()
    sheet_names = list_sheet_names(file_name, file_bytes)
    if not sheet_names:
        raise SpreadsheetValidationError("文件中没有可读取的 Sheet。")
    selected = str(sheet_name or sheet_names[0])
    if selected not in sheet_names:
        raise SpreadsheetValidationError("选择的 Sheet 不存在。")

    if extension == ".xlsx":
        rows = _read_xlsx(file_bytes, selected, max_rows)
    elif extension == ".xls":
        rows = _read_xls(file_bytes, selected, max_rows)
    else:
        rows = _read_csv(file_bytes, max_rows)
    return SheetData(
        sheet_names=sheet_names,
        selected_sheet=selected,
        rows=tuple(rows),
    )


def _read_xlsx(
    file_bytes: bytes,
    sheet_name: str,
    max_rows: int | None,
) -> list[tuple[str, ...]]:
    workbook = _load_xlsx(file_bytes)
    try:
        worksheet = workbook[sheet_name]
        # Some third-party exporters incorrectly declare the used range as A1
        # even when the worksheet XML contains many rows and columns.
        worksheet.reset_dimensions()
        rows: list[tuple[str, ...]] = []
        for row_index, cells in enumerate(worksheet.iter_rows(), start=1):
            if max_rows is not None and row_index > max_rows:
                break
            values = tuple(
                _string_value(cell.value, number_format=cell.number_format)
                for cell in cells
            )
            rows.append(_trim_trailing_empty(values))
        return rows
    finally:
        workbook.close()


def _read_xls(
    file_bytes: bytes,
    sheet_name: str,
    max_rows: int | None,
) -> list[tuple[str, ...]]:
    workbook = xlrd.open_workbook(
        file_contents=file_bytes,
        on_demand=True,
        formatting_info=True,
    )
    try:
        worksheet = workbook.sheet_by_name(sheet_name)
        row_limit = worksheet.nrows if max_rows is None else min(worksheet.nrows, max_rows)
        rows: list[tuple[str, ...]] = []
        for row_index in range(row_limit):
            values: list[str] = []
            for column_index in range(worksheet.ncols):
                cell = worksheet.cell(row_index, column_index)
                number_format = ""
                if cell.xf_index is not None and workbook.xf_list:
                    xf = workbook.xf_list[cell.xf_index]
                    cell_format = workbook.format_map.get(xf.format_key)
                    if cell_format is not None:
                        number_format = cell_format.format_str
                values.append(
                    _string_value(cell.value, number_format=number_format)
                )
            rows.append(_trim_trailing_empty(tuple(values)))
        return rows
    finally:
        workbook.release_resources()


def _read_csv(
    file_bytes: bytes,
    max_rows: int | None,
) -> list[tuple[str, ...]]:
    text = _decode_csv(file_bytes)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    rows = []
    for row_index, row in enumerate(reader, start=1):
        if max_rows is not None and row_index > max_rows:
            break
        rows.append(_trim_trailing_empty(tuple(str(value) for value in row)))
    return rows


def _trim_trailing_empty(values: tuple[str, ...]) -> tuple[str, ...]:
    end = len(values)
    while end and values[end - 1] == "":
        end -= 1
    return values[:end]
